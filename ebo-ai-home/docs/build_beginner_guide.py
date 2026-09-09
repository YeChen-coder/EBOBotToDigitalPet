from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
ASSETS = ROOT / "docs" / "assets"
OUTPUT = ROOT / "docs" / "EBO_AI_新手使用手册.docx"
ERROR_IMAGE = ASSETS / "private-error-example.png"  # Supply your own sanitized image.

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "12324A"
MUTED = "586F7F"
LIGHT_BLUE = "E8EEF5"
LIGHT_GREEN = "EAF6EF"
LIGHT_GOLD = "FFF3D8"
LIGHT_RED = "FDECEC"
WHITE = "FFFFFF"
GRID = "AFC0CA"


def rgb(hex_value: str) -> RGBColor:
    return RGBColor.from_string(hex_value)


def set_run_font(run, name="Calibri", size=11, color=INK, bold=False, italic=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    run.font.size = Pt(size)
    run.font.color.rgb = rgb(color)
    run.bold = bold
    run.italic = italic


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[min(idx, len(widths_dxa) - 1)])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color=GRID, size="8"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), size)
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def format_cell_text(cell, bold=False, color=INK, size=10.5):
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.15
        for run in paragraph.runs:
            set_run_font(run, size=size, color=color, bold=bold)


def add_table(doc, headers, rows, widths, header_fill=LIGHT_BLUE):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        set_cell_shading(cell, header_fill)
        format_cell_text(cell, bold=True, color=DARK_BLUE)
    set_repeat_table_header(table.rows[0])
    set_row_cant_split(table.rows[0])
    for row_values in rows:
        cells = table.add_row().cells
        set_row_cant_split(table.rows[-1])
        for i, value in enumerate(row_values):
            cells[i].text = value
            format_cell_text(cells[i])
    set_table_geometry(table, widths)
    return table


def add_callout(doc, title, body, fill=LIGHT_BLUE, icon="要点"):
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [1700, 7660])
    set_table_borders(table, color=fill, size="10")
    left, right = table.rows[0].cells
    set_row_cant_split(table.rows[0])
    set_cell_shading(left, fill)
    set_cell_shading(right, fill)
    left.text = icon
    right.text = ""
    p = right.paragraphs[0]
    r = p.add_run(title)
    set_run_font(r, size=11.5, color=DARK_BLUE, bold=True)
    p2 = right.add_paragraph(body)
    p2.paragraph_format.space_after = Pt(0)
    for run in p2.runs:
        set_run_font(run, size=10.5)
    format_cell_text(left, bold=True, color=DARK_BLUE)
    return table


def add_para(doc, text="", *, size=11, color=INK, bold=False, italic=False,
             align=WD_ALIGN_PARAGRAPH.LEFT, before=0, after=6, line=1.25, keep=False):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.add_run(text)
    return p


def set_alt_text(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_image(doc, path, width, caption, alt):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    shape = p.add_run().add_picture(str(path), width=Inches(width))
    set_alt_text(shape, caption, alt)
    cap = add_para(doc, caption, size=9, color=MUTED, italic=True,
                   align=WD_ALIGN_PARAGRAPH.CENTER, after=8, line=1.0)
    cap.paragraph_format.keep_with_next = False
    return shape


def add_page_break(doc):
    doc.add_page_break()


def add_step(doc, number, title, body, fill):
    table = doc.add_table(rows=1, cols=2)
    set_table_geometry(table, [1100, 8260])
    set_table_borders(table, color=fill, size="8")
    left, right = table.rows[0].cells
    set_row_cant_split(table.rows[0])
    set_cell_shading(left, fill)
    set_cell_shading(right, "FFFFFF")
    left.text = str(number)
    p = left.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    format_cell_text(left, bold=True, color=DARK_BLUE, size=18)
    right.text = ""
    rp = right.paragraphs[0]
    rr = rp.add_run(title)
    set_run_font(rr, size=12, color=DARK_BLUE, bold=True)
    rb = right.add_paragraph(body)
    rb.paragraph_format.space_after = Pt(0)
    rb.paragraph_format.line_spacing = 1.2
    for run in rb.runs:
        set_run_font(run, size=10.5)
    add_para(doc, "", after=2)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1.0)
section.right_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.left_margin = Inches(1.0)
section.header_distance = Inches(0.492)
section.footer_distance = Inches(0.492)

# compact_reference_guide preset
styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
normal.font.size = Pt(11)
normal.font.color.rgb = rgb(INK)
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

heading_tokens = {
    1: (16, BLUE, 18, 10),
    2: (13, BLUE, 14, 7),
    3: (12, DARK_BLUE, 10, 5),
}
for level, (size, color, before, after) in heading_tokens.items():
    style = styles[f"Heading {level}"]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = rgb(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

# Quiet running header/footer used by the editorial_cover pattern.
header = section.header
hp = header.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
hr = hp.add_run("EBO AI 小伙伴｜新手使用手册")
set_run_font(hr, size=8.5, color=MUTED, bold=True)
footer = section.footer
fp = footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
fr = fp.add_run("第 ")
set_run_font(fr, size=8.5, color=MUTED)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
fp._p.append(field)
fr2 = fp.add_run(" 页")
set_run_font(fr2, size=8.5, color=MUTED)

# Cover: editorial_cover
add_para(doc, "新手操作手册", size=11, color="B07A18", bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, before=58, after=18)
add_para(doc, "EBO AI 小伙伴", size=30, color=INK, bold=True,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=8, line=1.0)
add_para(doc, "从“能看到画面”到“让 AI 看懂画面”",
         size=15, color=DARK_BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=28)
add_para(doc, "适用设备：EBO Air 2S｜平台：Home Assistant｜视觉模型：OpenAI",
         size=10.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=42)
add_image(doc, ASSETS / "architecture.png", 6.25,
          "系统全景：设备可以换，camera.* 接口保持稳定",
          "架构图：EBO 或普通摄像头经过适配层成为 Home Assistant camera 实体，再由 LLM Vision 调用 OpenAI。")
add_para(doc, "版本：2026-08-26｜为当前本机隔离环境编写",
         size=9.5, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, before=10, after=0)

add_page_break(doc)
add_heading(doc, "1. 先记住这三句话", 1)
add_callout(doc, "第一句：这次报错不是 LLM Vision 造成的",
            "LLM Vision 只在你按“AI 看一眼”时抓取一张图片。它不会持续占用直播，也不会让 WebRTC 失效。",
            fill=LIGHT_GREEN, icon="✓")
add_para(doc, "")
add_callout(doc, "第二句：EBO 休眠时，RTSP 视频地址会暂时不存在",
            "EBO 连续 5 分钟没有收到控制命令后，会退出视频会话去休眠。此时 RTSP 返回 404 是“没有视频源”，不是账号坏了。",
            fill=LIGHT_GOLD, icon="💤")
add_para(doc, "")
add_callout(doc, "第三句：现在的面板先唤醒，再按需打开直播",
            "休眠时不加载视频卡片。按“1. 唤醒画面”后等待约 8 秒会显示当前画面；点画面时才打开 WebRTC 直播。",
            fill=LIGHT_BLUE, icon="1→2")
add_heading(doc, "四个名字分别做什么", 2)
add_table(doc, ["名字", "最简单的理解"], [
    ("EBO / 普通摄像头", "负责把真实世界的画面送出来。"),
    ("Home Assistant", "总控制台：放面板、按钮、脚本、自动化和状态。"),
    ("LLM Vision", "Home Assistant 里的“图片转交员”，负责拿图并调用模型。"),
    ("OpenAI", "真正理解图片并返回中文描述的视觉模型。"),
], [2700, 6660])

add_page_break(doc)
add_heading(doc, "2. 你看到的 404 报错是什么意思", 1)
add_image(doc, ERROR_IMAGE, 3.15,
          "修复前：面板在 EBO 休眠时仍强制启动 WebRTC",
          "用户提供的错误截图，显示 WebRTC DESCRIBE 请求从 RTSP 地址得到 404 Not Found。")
add_table(doc, ["报错片段", "翻成大白话"], [
    ("DESCRIBE failed: 404", "Home Assistant 去问 RTSP：有视频吗？当前回答是没有。"),
    ("rtsp://…/ebo", "这是 EBO 引擎临时提供的直播地址。"),
    ("Failed to start WebRTC", "网页想把 RTSP 转成浏览器能播的 WebRTC，但源头正在休眠。"),
    ("Unavailable", "这一刻没有直播，不等于设备永远不可用。"),
], [3000, 6360], header_fill=LIGHT_RED)
add_heading(doc, "真正的时间顺序", 2)
add_para(doc, "EBO 正常直播 → 连续 5 分钟没有控制命令 → EBO 引擎退出会话 → RTSP 地址返回 404 → 旧面板继续强制请求 WebRTC → 出现红色错误。")
add_callout(doc, "为什么保留 5 分钟休眠？",
            "减少摄像头长期打开和电量消耗，也保留明确的隐私边界。",
            fill=LIGHT_GREEN, icon="安全")

add_page_break(doc)
add_heading(doc, "3. 修复后，最简单的使用方法", 1)
add_image(doc, ASSETS / "usage-flow.png", 6.35,
          "日常使用只需要记住：唤醒 → 看画面 → AI 看一眼",
          "流程图：EBO 从休眠开始，先唤醒画面，再查看直播，再调用 OpenAI，最后可在五分钟后自动休眠。")
add_step(doc, 1, "打开左侧“AI 小伙伴”", "进入“看看家里”页面。休眠时会看到文字说明，不会看到红色 404。", LIGHT_BLUE)
add_step(doc, 2, "想看画面：点“1. 唤醒画面”", "等大约 8 秒。摄像头卡片会出现；点画面可打开实时直播。这个按钮不调用 OpenAI。", LIGHT_GREEN)
add_step(doc, 3, "想让 AI 解释：点“2. AI 看一眼”", "它会先唤醒 EBO，再抓取一张画面交给 OpenAI。返回结果显示在“AI 看到的画面”。", "EFE8FF")
add_step(doc, 4, "摄像头已经醒着时", "可以点“摄像头已经醒着：只做 AI 分析”，少等一次唤醒时间。", LIGHT_GOLD)

add_page_break(doc)
add_heading(doc, "4. 当前面板长什么样", 1)
add_image(doc, ASSETS / "dashboard-verified.png", 6.0,
          "修复并验证后的 AI 小伙伴面板",
          "Home Assistant AI 小伙伴面板截图：包含唤醒、AI 分析、实时画面、分析结果、摄像头来源和 EBO 安全状态。")
add_table(doc, ["面板区域", "你该怎么用"], [
    ("1. 唤醒画面", "只连接 EBO 和直播，不花 OpenAI 调用费用。"),
    ("EBO 画面卡片", "默认显示当前画面；点卡片才打开实时直播，避免后台长期占用 WebRTC。"),
    ("2. AI 看一眼", "唤醒 + 等 8 秒 + 抓图 + OpenAI 中文描述。"),
    ("AI 看到的画面", "显示最近一次模型返回；不是实时字幕。"),
    ("摄像头来源", "这是可替换接口。以后可切到普通 camera.* 实体。"),
    ("AI 提示词", "直接改文字，就能改变 AI 关注点，不用改代码。"),
    ("立即停止", "机器人运动异常时优先按它；视觉分析本身不驱动机器人。"),
], [2700, 6660])

add_page_break(doc)
add_heading(doc, "5. “脚本”是什么：把多步操作做成一个按钮", 1)
add_para(doc, "脚本可以理解成一张“操作配方”。你按一次，它会按顺序完成多步动作。当前系统有三张配方：")
add_image(doc, ASSETS / "scripts-page.png", 6.45,
          "设置 → Automations & scenes → Scripts：当前有 3 个脚本",
          "Home Assistant 脚本列表截图，显示通用 AI 分析、EBO 唤醒摄像头和 EBO 唤醒后分析三个脚本。")
add_table(doc, ["脚本实体", "做什么", "是否绑定 EBO"], [
    ("script.ebo_wake_camera", "唤醒 EBO，等待 8 秒。", "是"),
    ("script.ai_camera_describe", "分析 input_select 里选中的 camera.*。", "否"),
    ("script.ebo_ai_look_and_describe", "先调用唤醒脚本，再调用通用 AI 分析脚本。", "只有外层是"),
], [3250, 4300, 1810])
add_callout(doc, "为什么拆成三张配方？",
            "因为“唤醒 EBO”是品牌专属动作，而“分析任意 camera.*”是通用动作。以后换普通摄像头，只需绕开第一张配方。",
            fill=LIGHT_GREEN, icon="解耦")

add_page_break(doc)
add_heading(doc, "6. 不写代码，也能做自己的自动化", 1)
add_image(doc, ASSETS / "automation-entry.png", 5.55,
          "设置 → Automations & scenes → Automations → Create automation",
          "Home Assistant 自动化入口截图，右下角有 Create automation 按钮。")
add_heading(doc, "例子：每天上午 10 点让 AI 看一眼", 2)
add_step(doc, 1, "点 Create automation", "选择 Create new automation。", LIGHT_BLUE)
add_step(doc, 2, "When（触发条件）", "选择 Time，时间设成 10:00:00。", LIGHT_GREEN)
add_step(doc, 3, "Then do（执行动作）并保存", "动作填 script.ebo_ai_look_and_describe。保存后用 Run actions 测试；建议先从每天 1 次开始，因为每次会把一张图片发送给 OpenAI。", "EFE8FF")

add_page_break(doc)
add_heading(doc, "7. 改提示词，就是最简单的“编程”", 1)
add_para(doc, "在面板的“AI 提示词”输入框里直接改文字。下一次分析会使用新文字；不需要重启 Home Assistant。")
add_table(doc, ["想实现什么", "可以填写的提示词"], [
    ("陪伴式描述", "请用温暖、简短的中文描述画面，并提出一个轻松、不冒犯的小话题。"),
    ("找常见物品", "请列出画面中清楚可见的常见物品；不确定的不要猜。"),
    ("看宠物", "请只描述是否看见宠物、宠物大概在做什么；不要推断健康状况。"),
    ("看房间变化", "请简要描述房间里最明显的三处物体或布局；只说可见事实。"),
], [2600, 6760])
add_heading(doc, "给会写一点 YAML 的你：最小调用", 2)
code = add_callout(doc, "在自动化动作里调用现成脚本",
                   "action: script.ebo_ai_look_and_describe",
                   fill="F2F4F7", icon="YAML")
for cell in code.rows[0].cells:
    for p in cell.paragraphs:
        for run in p.runs:
            if "action:" in run.text:
                set_run_font(run, name="Consolas", size=10.5, color=INK)
add_para(doc, "自动化只调用这个稳定入口；OpenAI 的复杂细节继续留在通用脚本里。")

add_page_break(doc)
add_heading(doc, "8. 以后换普通摄像头，哪些东西需要改", 1)
add_image(doc, ASSETS / "architecture.png", 6.45,
          "解耦架构：只替换摄像头来源和适配层",
          "架构图强调 camera.* 是稳定接口，EBO 或普通摄像头均可接入 Home Assistant 和 LLM Vision。")
add_table(doc, ["部分", "换摄像头后"], [
    ("摄像头设备", "替换：EBO 可换成 RTSP、ONVIF 或其他 Home Assistant 摄像头。"),
    ("摄像头适配层", "替换：安装对应集成，让它产生一个 camera.xxx 实体。"),
    ("input_select.ai_camera_source", "只需把新 camera.xxx 加到选项并选中。"),
    ("script.ai_camera_describe", "不改。它只认 camera.*。"),
    ("LLM Vision / OpenAI", "不改。继续从 camera.* 抓图。"),
    ("Dashboard / 提示词", "原则上不改；如不再需要 EBO，只隐藏 EBO 专属按钮。"),
], [3000, 6360])
add_page_break(doc)
add_heading(doc, "接入普通摄像头的四步", 2)
add_step(doc, 1, "在 Home Assistant 添加摄像头集成", "常见选择是 ONVIF 或 Generic Camera，取决于摄像头提供什么地址。", LIGHT_BLUE)
add_step(doc, 2, "确认出现 camera.xxx", "先在开发者工具或设备页面确认能取到一张图。", LIGHT_GREEN)
add_step(doc, 3, "把 camera.xxx 加到摄像头来源", "修改 input_select.ai_camera_source 的 options，然后选择它。", "EFE8FF")
add_step(doc, 4, "点“只做 AI 分析”测试", "能返回描述，就说明更换完成。", LIGHT_GOLD)

add_page_break(doc)
add_heading(doc, "9. 出问题时，先看这张表", 1)
add_table(doc, ["现象", "最可能原因", "先做什么"], [
    ("休眠提示，没有画面", "EBO 正常休眠。", "点“1. 唤醒画面”，等 8 秒。"),
    ("仍出现 RTSP 404", "页面缓存了旧面板，或唤醒还没完成。", "刷新页面，再点唤醒；等待 10 秒。"),
    ("直播有画面，AI 不返回", "OpenAI Key、余额、网络或模型调用失败。", "看 设置 → 系统 → 日志，搜索 llmvision。"),
    ("AI 返回“分析失败”", "抓图或模型调用失败。", "先确认直播能看，再重试一次。"),
    ("EBO 离线", "机器人、家庭网络或 Enabot 云连接不可用。", "先在手机 EBO App 中确认设备在线。"),
    ("换摄像头后分析旧画面", "input_select 还选着 EBO。", "在“摄像头来源”选择新的 camera.xxx。"),
], [2400, 3300, 3660], header_fill=LIGHT_RED)
add_heading(doc, "查看日志的位置", 2)
add_para(doc, "Home Assistant 左下角 Settings → System → Logs。常用关键词：ebo、camera、rtsp、webrtc、llmvision、openai。")
add_callout(doc, "看到 404 时先判断 EBO 是否休眠",
            "如果刚好超过 5 分钟没有操作，先按唤醒，不要马上重装 LLM Vision 或删除集成。",
            fill=LIGHT_GOLD, icon="先判断")

add_page_break(doc)
add_heading(doc, "10. 隐私、安全与当前文件位置", 1)
add_heading(doc, "建议的使用边界", 2)
add_table(doc, ["建议", "原因"], [
    ("AI 只负责“看”和“描述”", "暂时不要让视觉模型直接驱动 EBO 移动；把运动动作放在人工确认之后。"),
    ("保留 5 分钟休眠", "减少长期视频会话、耗电和不必要的家庭画面传输。"),
    ("提示词要求只说可见事实", "避免猜测身份、健康、地址和其他敏感信息。"),
    ("OpenAI Key 只放 Home Assistant 配置", "不要写进 Dashboard、截图、文档或 Git 仓库。"),
], [3000, 6360])
add_heading(doc, "当前隔离项目的位置", 2)
add_table(doc, ["内容", "位置"], [
    ("整个项目", str(ROOT)),
    ("AI 通用脚本", r"homeassistant-config\packages\ai_camera_adapter.yaml"),
    ("AI 小伙伴面板", r"homeassistant-config\dashboards\ai-playmate.yaml"),
    ("本手册", r"docs\EBO_AI_新手使用手册.docx"),
], [2500, 6860])
add_callout(doc, "隔离性说明",
            "Home Assistant、EBO 引擎和配置都放在 ebo-ai-home 目录及其 Docker Compose 容器里。后续改动继续限定在这个项目，避免污染其他项目。",
            fill=LIGHT_GREEN, icon="隔离")
add_page_break(doc)
add_heading(doc, "一页速记：以后忘了就看这里", 1)
add_image(doc, ASSETS / "usage-flow.png", 6.35,
          "速记流程：唤醒 → 看画面 → AI 看一眼",
          "EBO AI 小伙伴速记流程图：先唤醒，再看实时画面，最后按需调用 OpenAI。")
add_callout(doc, "没有画面时",
            "先点“1. 唤醒画面”，等待约 8 秒。超过 5 分钟没操作后再次休眠是正常的。",
            fill=LIGHT_GOLD, icon="1")
add_para(doc, "")
add_callout(doc, "想要 AI 描述时",
            "点“2. AI 看一眼”。LLM Vision 负责取图，OpenAI 负责返回中文；它们不会抢走直播。",
            fill="EFE8FF", icon="2")
add_heading(doc, "最后的心智模型", 2)
add_para(doc, "把它想成插座：摄像头插进 camera.*；Home Assistant 决定什么时候取图；LLM Vision 把图交给 OpenAI；OpenAI 只返回文字。以后换摄像头，只换插头，不重做整间房。",
         size=12, color=DARK_BLUE, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=8, after=4)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
